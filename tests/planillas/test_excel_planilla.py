"""Bloque de firmas del Excel oficial: 4 en II.EE. y 3 en Sede."""
import io
import sys
import unittest

sys.path.insert(0, '/app')

from openpyxl import load_workbook

from cafae.documentos.excel_planillas import generar_excel_planilla

LINEA_FIRMA = '___________________________________'


def _info(tipo):
    return {'tipo': tipo, 'anio': 2026, 'mes': 3, 'numero': 1}


def _filas():
    return [{
        'dni': '12345678',
        'apellidos': 'PEREZ',
        'nombres': 'JUAN',
        'importe': 1000,
        'nivel_nombre': 'PRIMARIA',
    }]


def _firmas(cantidad):
    return [
        {'nombre_completo': f'Autoridad {i}', 'cargo': f'Cargo {i}'}
        for i in range(1, cantidad + 1)
    ]


def _hoja(tipo, firmas):
    datos = generar_excel_planilla(_info(tipo), _filas(), firmas)
    return load_workbook(io.BytesIO(datos)).active


def _lineas_de_firma(hoja):
    return [
        celda for fila in hoja.iter_rows() for celda in fila
        if celda.value == LINEA_FIRMA
    ]


class BloqueDeFirmas(unittest.TestCase):
    def test_iiee_dibuja_cuatro_firmas(self):
        hoja = _hoja('IIEE', _firmas(4))
        self.assertEqual(len(_lineas_de_firma(hoja)), 4)
        nombres = {c.value for fila in hoja.iter_rows() for c in fila}
        for i in range(1, 5):
            self.assertIn(f'Autoridad {i}', nombres)

    def test_sede_dibuja_tres_firmas(self):
        """La cuarta posición era un espacio vacío con un cargo que no existe en Sede."""
        hoja = _hoja('SEDE', _firmas(3))
        self.assertEqual(len(_lineas_de_firma(hoja)), 3)
        nombres = {c.value for fila in hoja.iter_rows() for c in fila}
        self.assertIn('Autoridad 3', nombres)
        self.assertNotIn('Autoridad 4', nombres)
        self.assertNotIn('Especialista en Finanzas', nombres)

    def test_sede_sin_firmas_usa_los_tres_cargos_de_la_sede(self):
        hoja = _hoja('SEDE', [])
        cargos = {c.value for fila in hoja.iter_rows() for c in fila}
        self.assertIn('Director de UGEL CHEPEN', cargos)
        self.assertIn('Directora Gestión Institucional', cargos)
        self.assertIn('Director de Gestión Administrativa', cargos)
        self.assertNotIn('Especialista en Finanzas', cargos)


if __name__ == '__main__':
    unittest.main()
