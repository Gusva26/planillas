"""Inspección de solo lectura de las referencias; nunca guarda los libros."""
import hashlib
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.formula.translate import Translator

ROOT = Path(__file__).resolve().parents[1]
FILES = {
    'sede': Path('C:/Users/Gusva/Downloads/PLANILLA CAFAE - SEDE UGELCH_SETIEMBRE_BD SISTEMAS.xlsx'),
    'iiee': Path('C:/Users/Gusva/Downloads/PLANILLA CAFAE - II.EE.s DSETIEMBRE 2026_SGD BD _SISTEMAS.xlsx'),
    'remuneraciones': Path('C:/Users/Gusva/Downloads/HOJA DE TRABAJO DE REMUNERACIONES - PLANILLA  CAFAE- 2026.xlsx'),
    'boleta': Path('C:/Users/Gusva/Downloads/GENERAR BOLETA CAFE.xlsx'),
}

def value(v):
    return v if v is None or isinstance(v, (str, int, float, bool)) else str(v)

def inventory():
    result = {}
    for key, path in FILES.items():
        wb = openpyxl.load_workbook(path, data_only=False)
        cached = openpyxl.load_workbook(path, data_only=True)
        out = {'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
               'calculation': str(wb.calculation), 'names': {n: str(d.attr_text) for n, d in wb.defined_names.items()},
               'external_links': len(wb._external_links), 'sheets': []}
        with zipfile.ZipFile(path) as z:
            out['macros'] = any('vbaProject' in n for n in z.namelist())
            out['media'] = [n for n in z.namelist() if n.startswith('xl/media/')]
        for ws in wb:
            formulas = [c for c in ws._cells.values() if c.data_type == 'f']
            errors = [[c.coordinate, c.value] for c in cached[ws.title]._cells.values() if c.data_type == 'e']
            patterns = {}
            for c in formulas:
                formula = str(c.value)
                # Relative structure only; original examples remain locatable by address.
                pattern = re.sub(r'(\$?[A-Z]{1,3}\$?)\d+', r'\1#', formula)
                entry = patterns.setdefault(pattern, {'count': 0, 'cell': c.coordinate, 'formula': formula})
                entry['count'] += 1
            out['sheets'].append({
                'name': ws.title, 'state': ws.sheet_state, 'dimensions': ws.calculate_dimension(),
                'print_area': str(ws.print_area), 'print_titles': ws.print_title_rows,
                'orientation': ws.page_setup.orientation, 'paper': ws.page_setup.paperSize,
                'scale': ws.page_setup.scale, 'fit_width': ws.page_setup.fitToWidth,
                'fit_height': ws.page_setup.fitToHeight, 'margins': str(ws.page_margins),
                'row_breaks': [b.id for b in ws.row_breaks.brk],
                'col_breaks': [b.id for b in ws.col_breaks.brk],
                'freeze': str(ws.freeze_panes), 'protected': ws.protection.sheet,
                'tables': list(ws.tables), 'validation_count': len(ws.data_validations.dataValidation),
                'images': len(ws._images), 'merges': [str(x) for x in ws.merged_cells.ranges],
                'formulas_count': len(formulas), 'formula_patterns': list(patterns.values()),
                'cached_error_count': len(errors), 'cached_errors': errors[:25],
                'header_cells': {c.coordinate: ('[dato personal omitido]' if key == 'boleta' and c.coordinate == 'I2' else value(c.value)) for c in ws._cells.values() if c.row <= (4 if ws.title == 'ADICIONAL AGOSTO' else 6) and c.value is not None},
            })
        result[key] = out
    dest = ROOT / 'analysis' / 'referencias'
    dest.mkdir(parents=True, exist_ok=True)
    (dest / 'inventario.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    for key, book in result.items():
        print(json.dumps({'book': key, **{k:v for k,v in book.items() if k != 'sheets'}}, ensure_ascii=False))
        for s in book['sheets']:
            print(json.dumps({k:v for k,v in s.items() if k not in ('merges','formula_patterns','margins','header_cells')}, ensure_ascii=False))
            print('HEADERS', json.dumps(s['header_cells'], ensure_ascii=False))
            print('FORMULAS', json.dumps(s['formula_patterns'][:20], ensure_ascii=False))

def detail():
    targets = {
        'sede': [('PLANILLA SEDE SET_26', [(35, 56)])],
        'iiee': [('PLANILLA SET_26', [(5,11),(157,188)]), ('hoja de trabajo', [(6,18)])],
        'remuneraciones': [('CAFAE - JULIO', [(7,24)])],
        'boleta': [('ENE26', [(17,37)])],
    }
    for key, sheets in targets.items():
        wb = openpyxl.load_workbook(FILES[key], data_only=False)
        cache = openpyxl.load_workbook(FILES[key], data_only=True)
        for name, spans in sheets:
            ws = wb[name]
            print('\nBOOK', key, 'SHEET', name)
            for low, high in spans:
                for r in range(low, high+1):
                    cells = []
                    for c in ws._cells.values():
                        if c.row != r or c.value is None:
                            continue
                        if c.column > 45:
                            continue
                        v = value(c.value)
                        cells.append([c.coordinate, v, value(cache[name][c.coordinate].value)] if c.data_type=='f' else [c.coordinate,v])
                    if cells:
                        print(json.dumps(cells, ensure_ascii=False))

def diagnostics():
    for key, name, dni_col, money_cols, low, high in [
        ('iiee','PLANILLA SET_26','B',list('PQRSTUVWXYZ')+['AA'],12,156),
        ('sede','PLANILLA SEDE SET_26','G',list('HIJKLMNOPQ'),7,35),
    ]:
        wb = openpyxl.load_workbook(FILES[key], data_only=False)
        ws = wb[name]
        cache = openpyxl.load_workbook(FILES[key], data_only=True)[name]
        identified=[]; numbered=[]; missing=[]; ids=[]; nonzero=[]
        for r in range(low,high+1):
            d=cache[f'{dni_col}{r}'].value
            item=cache[f'A{r}'].value
            if isinstance(item,(int,float)):
                numbered.append(r)
                if d is None: missing.append(r)
            if d is not None and re.fullmatch(r'\d{7,8}',str(d)):
                identified.append(r); ids.append(str(d))
            monies={f'{col}{r}': {'value':value(cache[f'{col}{r}'].value), 'formula': ws[f'{col}{r}'].value if ws[f'{col}{r}'].data_type=='f' else None} for col in money_cols if cache[f'{col}{r}'].value not in (None,0)}
            if monies:
                nonzero.append({'row':r,'is_numbered':isinstance(item,(int,float)),'dni_present':d is not None,'cells':monies})
        print(json.dumps({'book':key,'numbered_rows':numbered,'rows_with_dni':identified,'numbered_missing_dni':missing,'duplicate_dni_groups':sum(1 for n in Counter(ids).values() if n>1),'nonzero_money':nonzero},ensure_ascii=False))
        print('hidden_columns', [(k,v.hidden,v.width) for k,v in ws.column_dimensions.items() if v.hidden])
        print('hidden_rows', [r for r,d in ws.row_dimensions.items() if d.hidden])
        print('footer_merges',[str(m) for m in ws.merged_cells.ranges if m.min_row > high])
    for key in ('iiee','remuneraciones','boleta','sede'):
        with zipfile.ZipFile(FILES[key]) as z:
            parts=[n for n in z.namelist() if re.search(r'(drawings|vml|externalLinks)',n)]
            print('DRAWINGS',key,parts)
    wb=openpyxl.load_workbook(FILES['boleta'],data_only=False)
    print('BOLETA_LABELS',json.dumps({c.coordinate: value(c.value) for c in wb.active._cells.values() if c.column in (1,2,6,8) and c.row<=23 and c.value is not None},ensure_ascii=False))
    w=openpyxl.load_workbook(FILES['iiee'],data_only=False)['GENERADOR TXT']
    v=openpyxl.load_workbook(FILES['iiee'],data_only=True)['GENERADOR TXT']
    print('TXT_AUDIT',json.dumps({f'{col}{r}':{'type':w[f'{col}{r}'].data_type,'format':w[f'{col}{r}'].number_format,'value':value(w[f'{col}{r}'].value),'cached':value(v[f'{col}{r}'].value),'cached_length':len(str(v[f'{col}{r}'].value)) if v[f'{col}{r}'].value is not None else None} for r in (3,4,5,139,140) for col in ('B','C','D','N','O','P','Q','R','T','U','V')},ensure_ascii=False))


def controls():
    out={}
    for key,name,dni,account,amount,disc,net,low,high in [
        ('iiee','PLANILLA SET_26','B','O','P',list('RSTUVWXY'),'AA',12,155),
        ('sede','PLANILLA SEDE SET_26','G','F','J',list('KLMNO'),'Q',7,34)]:
        w=openpyxl.load_workbook(FILES[key],data_only=False)[name]
        c=openpyxl.load_workbook(FILES[key],data_only=True)[name]
        rows=[r for r in range(low,high+1) if (isinstance(c[f'A{r}'].value,(int,float)) and (key=='iiee' or c[f'D{r}'].value is not None))]
        def dec(v):
            return Decimal(str(v)) if isinstance(v,(int,float)) else Decimal(0)
        amounts=sum((dec(c[f'{amount}{r}'].value) for r in rows),Decimal(0))
        discounts=sum((sum((dec(c[f'{col}{r}'].value) for col in disc),Decimal(0)) for r in rows),Decimal(0))
        out[key]={'detail_rows':len(rows),'dni_count':sum(c[f'{dni}{r}'].value is not None for r in rows),
             'dni_lengths':dict(Counter(len(str(c[f'{dni}{r}'].value)) for r in rows if c[f'{dni}{r}'].value is not None)),
             'account_count':sum(isinstance(c[f'{account}{r}'].value,(int,float)) or (isinstance(c[f'{account}{r}'].value,str) and c[f'{account}{r}'].value.isdigit()) for r in rows),
             'amount_numeric_rows':sum(isinstance(c[f'{amount}{r}'].value,(int,float)) for r in rows),
             'amount_sum':str(amounts),'discount_sum':str(discounts),'net_expected_for_populated_amounts':str(amounts-discounts),
             'net_numeric_rows':sum(isinstance(c[f'{net}{r}'].value,(int,float)) for r in rows),
             'net_cached_sum':str(sum((dec(c[f'{net}{r}'].value) for r in rows),Decimal(0))),
             'formulas_net_count':sum(w[f'{net}{r}'].data_type=='f' for r in rows),
             'specific_controls':{x:{'value':value(c[x].value),'formula':value(w[x].value)} for x in (['P157','Q157','P158','Z157','AA157','D179','Z14','AA14','AA15','AA16'] if key=='iiee' else ['H36','I36','J36','Q36'])},
             'group_labels':{cell.coordinate:value(cell.value) for cell in w._cells.values() if isinstance(cell.value,str) and not cell.value.startswith('=') and re.search(r'META|SUB.?TOTAL|TOTAL GIRADO',cell.value,re.I)}}
    for key in FILES:
        wb=openpyxl.load_workbook(FILES[key],data_only=False)
        deps={}
        for s in wb:
            deps[s.title]=[c.coordinate+': '+str(c.value) for c in s._cells.values() if c.data_type=='f' and ('!' in str(c.value) or '#REF!' in str(c.value))]
        out.setdefault(key,{})['cross_sheet_formulas']=deps
    dest=ROOT/'analysis'/'referencias'/'controles.json'
    dest.write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(out,ensure_ascii=False,indent=2))

def layout():
    for key,path in FILES.items():
        wb=openpyxl.load_workbook(path,data_only=False)
        print('BOOK',key)
        for s in wb:
            if s.sheet_state != 'visible':
                continue
            print('SHEET',s.title,'header_footer',str(s.HeaderFooter),'margins',str(s.page_margins))
            print('fonts',Counter((c.font.name,c.font.sz) for c in s._cells.values() if c.value is not None))
            print('hidden_rows',[r for r,v in s.row_dimensions.items() if v.hidden])
            print('hidden_cols',[k for k,v in s.column_dimensions.items() if v.hidden])
        with zipfile.ZipFile(path) as z:
            for part in z.namelist():
                if part.startswith('xl/drawings/') and part.endswith(('.xml','.vml')):
                    tree=ET.fromstring(z.read(part))
                    texts=[x.text for x in tree.iter() if x.text and x.text.strip() and x.tag.rsplit('}',1)[-1] in ('t','div','span','font')]
                    print(part,'texts',texts)
            if key=='boleta':
                (ROOT/'analysis'/'referencias'/'boleta_imagen_cabecera.jpeg').write_bytes(z.read('xl/media/image1.jpeg'))

def verify():
    path=ROOT/'analysis'/'referencias'/'inventario.json'
    manifest=json.loads(path.read_text(encoding='utf-8'))
    for key,source in FILES.items():
        same=hashlib.sha256(source.read_bytes()).hexdigest()==manifest[key]['sha256']
        print(key, 'ORIGINAL SIN CAMBIOS' if same else 'HASH DIFERENTE')
        if not same:
            raise RuntimeError('Cambió el hash del archivo original')
    # Reducir datos personales en el inventario estructural persistido.
    for s in manifest['sede']['sheets']:
        if s['name']=='ADICIONAL AGOSTO':
            s['header_cells']={k:v for k,v in s['header_cells'].items() if int(re.sub(r'\D','',k))<=4}
    manifest['boleta']['sheets'][0]['header_cells']['I2']='[dato personal omitido]'
    path.write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')

if __name__ == '__main__':
    {'detail':detail,'diagnostics':diagnostics,'inventory':inventory,'controls':controls,'layout':layout,'verify':verify}.get(sys.argv[1] if len(sys.argv)>1 else 'inventory', inventory)()
